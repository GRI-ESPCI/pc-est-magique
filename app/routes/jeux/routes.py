"""PC est magique - Jeux Routes"""

import datetime
import io
import logging
import os
import re
import threading
import time
import xml.sax.saxutils as saxutils
import zipfile

import defusedxml.ElementTree as ET
import flask
from flask_babel import _
import openpyxl
import sqlalchemy as sa
from sqlalchemy.orm import selectinload

from app import context, db
from app.models import (
    JeuItem,
    JeuReservation,
    Notification,
    PCeen,
    PermissionScope,
    PermissionType,
    PushSubscription,
)
from app.routes.jeux import bp
from app.routes.jeux.forms import AddJeuForm, UploadCatalogueForm
from app.utils import typing
from app.utils.holidays import compute_jeu_return_date


_ODS_NS = {
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}

_ALLOWED_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
_MAX_COVER_SIZE_BYTES = 10 * 1024 * 1024


def _normalize(s: str) -> str:
    """Lowercase and collapse whitespace for comparison."""
    return re.sub(r"\s+", " ", s.lower()).strip()


def _clean_str(val, max_len: int | None = None) -> str | None:
    """Convert a value to a stripped string (optionally truncated) or None."""
    if val is None:
        return None
    s = str(val).strip()
    if not s:
        return None
    return s[:max_len] if max_len is not None else s


def _redirect_to_main() -> typing.RouteReturn:
    """Redirect back to the Jeux catalogue preserving search and category filters."""
    search_q = _clean_str(flask.request.form.get("q"))
    selected_cat = _clean_str(flask.request.form.get("cat"))
    return flask.redirect(flask.url_for("jeux.main", q=search_q, cat=selected_cat))


def _flash_form_errors(form) -> None:
    """Flash all validation errors of a WTForm."""
    for errors in form.errors.values():
        for error in errors:
            flask.flash(f"{error}", "danger")


def _get_categories_with_counts() -> tuple[list[tuple[str, int]], int]:
    """Return sorted (categorie, count) pairs and total games count."""
    cat_rows = db.session.execute(
        sa.select(JeuItem.categorie, sa.func.count(JeuItem.id))
        .group_by(JeuItem.categorie)
        .order_by(sa.func.lower(JeuItem.categorie))
    ).all()
    categories = [(cat, count) for cat, count in cat_rows]
    return categories, sum(count for _, count in categories)


def _remove_cover_file(item: JeuItem) -> None:
    """Remove the uploaded cover file of a game from disk if present."""
    if not item.image_filename:
        return
    base_path = flask.current_app.config["JEUX_BASE_PATH"]
    old_path = os.path.join(base_path, item.image_filename)
    if os.path.exists(old_path):
        try:
            os.remove(old_path)
        except Exception as e:
            logging.error(f"Error removing game cover: {e}")
    item.image_filename = None


def _save_cover_file(item: JeuItem, cover_file) -> str | None:
    """Validate and save an uploaded cover file for a game.

    Returns:
        An error message string if validation failed, or None on success.
    """
    if not cover_file or not getattr(cover_file, "filename", None):
        return None

    ext = os.path.splitext(cover_file.filename.lower())[1]
    if ext not in _ALLOWED_IMAGE_EXTENSIONS:
        return str(_("Format d'image non supporté. Utilisez JPG, PNG, WEBP ou GIF."))

    cover_file.seek(0, os.SEEK_END)
    size = cover_file.tell()
    cover_file.seek(0)
    if size > _MAX_COVER_SIZE_BYTES:
        return str(_("L'image dépasse la taille maximale de 10 Mo."))

    base_path = flask.current_app.config["JEUX_BASE_PATH"]
    os.makedirs(base_path, exist_ok=True)
    _remove_cover_file(item)

    safe_name = f"jeu_{item.id}_{int(time.time())}{ext}"
    cover_file.save(os.path.join(base_path, safe_name))
    item.image_filename = safe_name
    return None


def _parse_ods_rows(file_stream) -> list[list[str]]:
    """Extract rows from all sheets of an ODS file using standard zipfile + XML."""
    rows: list[list[str]] = []
    with zipfile.ZipFile(file_stream) as zf:
        with zf.open("content.xml") as content_file:
            tree = ET.parse(content_file)
    root = tree.getroot()
    table_ns = _ODS_NS["table"]
    repeat_row_attr = f"{{{table_ns}}}number-rows-repeated"
    repeat_col_attr = f"{{{table_ns}}}number-columns-repeated"

    for table in root.findall(".//table:table", _ODS_NS):
        table_rows = table.findall("table:table-row", _ODS_NS)
        for row_idx, tr in enumerate(table_rows):
            if row_idx == 0:
                # Skip header row
                continue
            row_repeat = int(tr.get(repeat_row_attr, "1"))
            if row_repeat > 100:
                # Large trailing empty rows block in ODS
                continue

            cells: list[str] = []
            for tc in tr.findall("table:table-cell", _ODS_NS):
                col_repeat = min(int(tc.get(repeat_col_attr, "1")), 20)
                paragraphs = [
                    "".join(p.itertext()).strip()
                    for p in tc.findall("text:p", _ODS_NS)
                ]
                cell_text = "\n".join(p for p in paragraphs if p).strip()
                for _ in range(col_repeat):
                    cells.append(cell_text)
                    if len(cells) >= 10:
                        break
                if len(cells) >= 10:
                    break

            if any(cells):
                for _ in range(min(row_repeat, 20)):
                    rows.append(list(cells))

    return rows


def _parse_xlsx_rows(file_stream) -> list[list[str]]:
    """Extract rows from all sheets of an XLSX file."""
    wb = openpyxl.load_workbook(file_stream, read_only=True, data_only=True)
    rows: list[list[str]] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            cells = [str(c).strip() if c is not None else "" for c in row[:10]]
            if any(cells):
                rows.append(cells)
    wb.close()
    return rows


def _export_catalogue_ods(items: list[JeuItem]) -> io.BytesIO:
    """Generate an ODS file in memory containing the Jeux catalogue."""
    headers = [
        "Nom",
        "Catégorie",
        "Temps de jeu",
        "Nombre de joueurs",
        "Popularité",
        "Commentaires",
        "Image URL",
    ]

    def _cell_xml(val: str | None) -> str:
        if not val:
            return "<table:table-cell/>"
        lines = str(val).splitlines() or [""]
        paragraphs = "".join(f"<text:p>{saxutils.escape(line)}</text:p>" for line in lines)
        return f'<table:table-cell office:value-type="string">{paragraphs}</table:table-cell>'

    row_xmls: list[str] = []
    # Header row
    header_cells = "".join(_cell_xml(h) for h in headers)
    row_xmls.append(f"<table:table-row>{header_cells}</table:table-row>")

    for item in items:
        comment = item.nota_bene if item.nota_bene is not None else item.commentaires
        row_vals = [
            item.nom,
            item.categorie,
            item.temps_de_jeu,
            item.nb_joueurs,
            item.popularite,
            comment,
            item.image_url,
        ]
        cells_xml = "".join(_cell_xml(v) for v in row_vals)
        row_xmls.append(f"<table:table-row>{cells_xml}</table:table-row>")

    rows_str = "".join(row_xmls)
    content_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content'
        ' xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"'
        ' xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"'
        ' office:version="1.2">'
        "<office:body>"
        "<office:spreadsheet>"
        '<table:table table:name="Catalogue Jeux">'
        f"{rows_str}"
        "</table:table>"
        "</office:spreadsheet>"
        "</office:body>"
        "</office:document-content>"
    )

    manifest_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">'
        '<manifest:file-entry manifest:full-path="/" manifest:version="1.2" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>"
    )

    mimetype = "application/vnd.oasis.opendocument.spreadsheet"

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("mimetype", mimetype, compress_type=zipfile.ZIP_STORED)
        zf.writestr("META-INF/manifest.xml", manifest_xml.encode("utf-8"))
        zf.writestr("content.xml", content_xml.encode("utf-8"))

    buf.seek(0)
    return buf


def _import_catalogue(file_stream, filename: str) -> tuple[int, int]:
    """Parse an ODS or XLSX catalogue and upsert games into the database.

    Columns expected:
        A (0): Nom
        B (1): Catégorie
        C (2): Temps de jeu
        D (3): Nbre de joueurs
        E (4): Popularité
        F (5): Commentaires
        G (6): URL de l'image

    Returns:
        Tuple of (added_count, updated_count).
    """
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".ods":
        raw_rows = _parse_ods_rows(file_stream)
    else:
        raw_rows = _parse_xlsx_rows(file_stream)

    existing: dict[str, JeuItem] = {}
    for item in db.session.scalars(sa.select(JeuItem)).all():
        existing[_normalize(item.nom)] = item

    added = 0
    updated = 0

    for row in raw_rows:
        nom = _clean_str(row[0], 500) if len(row) > 0 else None
        if not nom or _normalize(nom) == "nom":
            continue

        categorie = (_clean_str(row[1], 100) if len(row) > 1 else None) or "Autres"
        temps_de_jeu = _clean_str(row[2], 100) if len(row) > 2 else None
        nb_joueurs = _clean_str(row[3], 100) if len(row) > 3 else None
        popularite = _clean_str(row[4], 100) if len(row) > 4 else None
        commentaires = _clean_str(row[5]) if len(row) > 5 else None
        image_url = _clean_str(row[6], 500) if len(row) > 6 else None

        key = _normalize(nom)
        if key not in existing:
            item = JeuItem(
                nom=nom,
                categorie=categorie,
                temps_de_jeu=temps_de_jeu,
                nb_joueurs=nb_joueurs,
                popularite=popularite,
                commentaires=commentaires,
                nota_bene=commentaires,
                image_url=image_url,
            )
            db.session.add(item)
            existing[key] = item
            added += 1
        else:
            item = existing[key]
            changed = False
            for attr, val in (
                ("categorie", categorie),
                ("temps_de_jeu", temps_de_jeu),
                ("nb_joueurs", nb_joueurs),
                ("popularite", popularite),
                ("image_url", image_url),
            ):
                if val is not None and getattr(item, attr) != val:
                    setattr(item, attr, val)
                    changed = True
            if commentaires is not None and item.commentaires != commentaires:
                item.commentaires = commentaires
                if not item.nota_bene:
                    item.nota_bene = commentaires
                changed = True
            if changed:
                updated += 1

    db.session.commit()
    return added, updated


def _notify_reservation_cancelled(pceen_id: int, jeu_nom: str) -> None:
    """Create an in-app and WebPush notification when a reservation is cancelled."""
    title = str(_("Réservation annulée"))
    body = str(
        _(
            "Votre réservation pour le jeu « %(nom)s » a été annulée par le club Jeux.",
            nom=jeu_nom,
        )
    )
    url = flask.url_for("jeux.main")

    notif = Notification(
        title=title[:128],
        body=body[:512],
        url=url,
        target_type="pceen",
        pceen_id=pceen_id,
    )
    db.session.add(notif)
    db.session.commit()

    subs = db.session.scalars(
        sa.select(PushSubscription).where(PushSubscription._pceen_id == pceen_id)
    ).all()
    if not subs:
        return

    vapid_private_key = flask.current_app.config.get("VAPID_PRIVATE_KEY_PATH")
    if not vapid_private_key or not os.path.exists(vapid_private_key):
        return

    subs_data = [
        {
            "id": s.id,
            "endpoint": s.endpoint,
            "p256dh": s.p256dh,
            "auth": s.auth,
        }
        for s in subs
    ]
    message = {
        "title": title,
        "body": body,
        "image": "",
        "url": url,
        "quiet": False,
        "id": notif.id,
    }
    vapid_claims = {"sub": "mailto:gri@espci.fr"}
    app = flask.current_app._get_current_object()

    from app.routes.gris.routes import _send_push_notifications

    threading.Thread(
        target=_send_push_notifications,
        args=(app, subs_data, message, vapid_private_key, vapid_claims),
    ).start()


@bp.route("")
@bp.route("/")
def main() -> typing.RouteReturn:
    """Jeux catalogue – user view."""
    search = flask.request.args.get("q", "").strip()
    selected_cat = flask.request.args.get("cat", "").strip()

    today = datetime.date.today()
    # Order by category, then by name inside a given category (case-insensitive)
    query = (
        sa.select(JeuItem)
        .options(
            selectinload(
                JeuItem.reservations.and_(
                    JeuReservation.cancelled.is_(False),
                    JeuReservation.end_date >= today,
                )
            ).selectinload(JeuReservation.pceen)
        )
        .order_by(
            sa.func.lower(JeuItem.categorie),
            sa.func.lower(JeuItem.nom),
        )
    )

    if selected_cat:
        query = query.where(JeuItem.categorie == selected_cat)

    if search:
        pattern = f"%{search}%"
        query = query.where(
            sa.or_(
                JeuItem.nom.ilike(pattern),
                JeuItem.categorie.ilike(pattern),
                JeuItem.temps_de_jeu.ilike(pattern),
                JeuItem.nb_joueurs.ilike(pattern),
                JeuItem.commentaires.ilike(pattern),
                JeuItem.nota_bene.ilike(pattern),
            )
        )

    items = db.session.scalars(query).all()
    categories, total_jeux = _get_categories_with_counts()

    can_admin = context.g.is_gri or (
        context.g.logged_in
        and context.g.pceen.has_permission(PermissionType.write, PermissionScope.jeux)
    )

    return flask.render_template(
        "jeux/main.html",
        title=_("Catalogue Jeux"),
        items=items,
        search=search,
        selected_cat=selected_cat,
        categories=categories,
        total_jeux=total_jeux,
        can_admin=can_admin,
        next_return_date=compute_jeu_return_date(today),
    )


@bp.route("/covers/<filename>")
def serve_cover(filename: str) -> typing.RouteReturn:
    """Serve uploaded game cover image."""
    return flask.send_from_directory(
        flask.current_app.config["JEUX_BASE_PATH"], filename
    )


@bp.route("/reserve/<int:item_id>", methods=["POST"])
@context.logged_in_only
def reserve_item(item_id: int) -> typing.RouteReturn:
    """Allow a logged-in PCeen to reserve an available game for a given duration."""
    item = db.session.scalar(
        sa.select(JeuItem).where(JeuItem.id == item_id).with_for_update()
    )
    if not item:
        flask.flash(_("Jeu introuvable."), "danger")
        return flask.redirect(flask.url_for("jeux.main"))

    today = datetime.date.today()
    already_reserved = (
        db.session.scalar(
            sa.select(JeuReservation.id)
            .where(
                JeuReservation._jeu_id == item.id,
                JeuReservation.cancelled.is_(False),
                JeuReservation.end_date >= today,
            )
            .limit(1)
        )
        is not None
    )
    if already_reserved:
        flask.flash(_("Ce jeu est déjà réservé."), "warning")
        return _redirect_to_main()

    end_date = compute_jeu_return_date(today)

    reservation = JeuReservation(
        _jeu_id=item.id,
        _pceen_id=context.g.pceen.id,
        start_date=today,
        end_date=end_date,
        cancelled=False,
    )
    db.session.add(reservation)
    db.session.commit()

    flask.flash(
        _(
            "Vous avez réservé « %(nom)s » jusqu'au %(date)s.",
            nom=item.nom,
            date=end_date.strftime("%d/%m/%Y"),
        ),
        "success",
    )
    return _redirect_to_main()


@bp.route("/reservation/<int:res_id>/cancel", methods=["POST"])
@context.permission_only(PermissionType.write, PermissionScope.jeux)
def cancel_reservation(res_id: int) -> typing.RouteReturn:
    """Allow a Jeux admin to cancel a reservation and notify the PCeen."""
    reservation = db.session.get(JeuReservation, res_id)
    from_admin = flask.request.form.get("from_admin") == "1"
    from_reservations = flask.request.form.get("from_reservations") == "1"

    def _redirect_after_cancel() -> typing.RouteReturn:
        if from_reservations:
            return flask.redirect(
                flask.url_for(
                    "jeux.reservations",
                    page=flask.request.form.get("page", 1, type=int),
                    sort=_clean_str(flask.request.form.get("sort")) or "date",
                    way=_clean_str(flask.request.form.get("way")) or "desc",
                    q=_clean_str(flask.request.form.get("q")),
                    status=_clean_str(flask.request.form.get("status")) or "all",
                )
            )
        if from_admin:
            return flask.redirect(flask.url_for("jeux.admin"))
        return _redirect_to_main()

    if not reservation:
        flask.flash(_("Réservation introuvable."), "danger")
        return _redirect_after_cancel()

    nom = reservation.jeu.nom if reservation.jeu else ""
    pceen_id = reservation._pceen_id
    pceen_name = reservation.pceen.full_name if reservation.pceen else ""
    reservation.cancelled = True
    db.session.commit()

    if pceen_id and nom:
        try:
            _notify_reservation_cancelled(pceen_id, nom)
        except Exception as e:
            logging.error(f"Error sending cancellation notification: {e}")

    flask.flash(
        _(
            "La réservation de « %(nom)s » par %(pceen)s a été annulée.",
            nom=nom,
            pceen=pceen_name,
        ),
        "success",
    )

    return _redirect_after_cancel()


@bp.route("/reservations", methods=["GET"])
@context.permission_only(PermissionType.write, PermissionScope.jeux)
def reservations() -> typing.RouteReturn:
    """Paginated reservation history for Jeux admins (similar to bar.transactions)."""
    page = flask.request.args.get("page", 1, type=int)
    sort = flask.request.args.get("sort", "date", type=str)
    way = flask.request.args.get("way", "desc", type=str)
    if way not in ("asc", "desc"):
        way = "desc"
    q = flask.request.args.get("q", "", type=str).strip()
    status = flask.request.args.get("status", "all", type=str).strip()
    if status not in ("all", "active", "completed", "cancelled"):
        status = "all"

    today = datetime.date.today()

    stmt = (
        sa.select(JeuReservation)
        .outerjoin(JeuItem, JeuReservation._jeu_id == JeuItem.id)
        .outerjoin(PCeen, JeuReservation._pceen_id == PCeen.id)
        .options(
            selectinload(JeuReservation.jeu),
            selectinload(JeuReservation.pceen),
        )
    )

    if status == "active":
        stmt = stmt.where(
            JeuReservation.cancelled.is_(False),
            JeuReservation.end_date >= today,
        )
    elif status == "completed":
        stmt = stmt.where(
            JeuReservation.cancelled.is_(False),
            JeuReservation.end_date < today,
        )
    elif status == "cancelled":
        stmt = stmt.where(JeuReservation.cancelled.is_(True))

    if q:
        pattern = f"%{q}%"
        search_filters = [
            JeuItem.nom.ilike(pattern),
            JeuItem.categorie.ilike(pattern),
            PCeen.nom.ilike(pattern),
            PCeen.prenom.ilike(pattern),
            (PCeen.prenom + " " + PCeen.nom).ilike(pattern),
            (PCeen.nom + " " + PCeen.prenom).ilike(pattern),
            PCeen.username.ilike(pattern),
        ]
        if q.isdigit():
            search_filters.append(PCeen.promo == int(q))
            search_filters.append(JeuReservation.id == int(q))
        stmt = stmt.where(sa.or_(*search_filters))

    if sort == "jeu":
        order_col = sa.func.lower(JeuItem.nom)
    elif sort == "pceen":
        order_col = sa.func.lower(PCeen.nom)
    elif sort == "start":
        order_col = JeuReservation.start_date
    elif sort == "end":
        order_col = JeuReservation.end_date
    else:
        sort = "date"
        order_col = JeuReservation.created_at

    if way == "asc":
        stmt = stmt.order_by(order_col.asc(), JeuReservation.id.asc())
    else:
        stmt = stmt.order_by(order_col.desc(), JeuReservation.id.desc())

    per_page = flask.current_app.config.get("BAR_ITEMS_PER_PAGE", 25) or 25
    paginator = db.paginate(stmt, page=page, per_page=per_page, error_out=False)

    return flask.render_template(
        "jeux/reservations.html",
        title=_("Historique des réservations – Jeux"),
        paginator=paginator,
        sort=sort,
        way=way,
        q=q,
        status=status,
        today=today,
    )


@bp.route("/admin", methods=["GET"])
@context.permission_only(PermissionType.write, PermissionScope.jeux)
def admin() -> typing.RouteReturn:
    """Jeux admin dashboard."""
    form = UploadCatalogueForm()
    add_form = AddJeuForm()

    categories, total_jeux = _get_categories_with_counts()

    today = datetime.date.today()
    active_reservations = db.session.scalars(
        sa.select(JeuReservation)
        .options(
            selectinload(JeuReservation.jeu),
            selectinload(JeuReservation.pceen),
        )
        .where(
            JeuReservation.cancelled.is_(False),
            JeuReservation.end_date >= today,
        )
        .order_by(JeuReservation.end_date.asc(), JeuReservation.id.desc())
    ).all()

    recent = db.session.scalars(
        sa.select(JeuItem).order_by(JeuItem.id.desc()).limit(10)
    ).all()

    return flask.render_template(
        "jeux/admin.html",
        title=_("Jeux – Administration"),
        form=form,
        add_form=add_form,
        categories=categories,
        total_jeux=total_jeux,
        active_reservations=active_reservations,
        recent=recent,
    )


@bp.route("/admin/add", methods=["POST"])
@context.permission_only(PermissionType.write, PermissionScope.jeux)
def admin_add() -> typing.RouteReturn:
    """Handle manual creation of a single game."""
    add_form = AddJeuForm()
    if not add_form.validate_on_submit():
        _flash_form_errors(add_form)
        return flask.redirect(flask.url_for("jeux.admin"))

    nom = _clean_str(add_form.nom.data, 500) or ""
    norm_nom = _normalize(nom)
    existing_match = db.session.scalar(
        sa.select(JeuItem).where(sa.func.lower(JeuItem.nom) == norm_nom)
    )
    if existing_match:
        flask.flash(
            _("Un jeu nommé « %(nom)s » existe déjà dans le catalogue.", nom=existing_match.nom),
            "warning",
        )
        return flask.redirect(flask.url_for("jeux.admin"))

    item = JeuItem(
        nom=nom,
        categorie=_clean_str(add_form.categorie.data, 100) or "Autres",
        temps_de_jeu=_clean_str(add_form.temps_de_jeu.data, 100),
        nb_joueurs=_clean_str(add_form.nb_joueurs.data, 100),
        popularite=_clean_str(add_form.popularite.data, 100),
        nota_bene=_clean_str(add_form.nota_bene.data),
        resume=_clean_str(add_form.resume.data),
        image_url=_clean_str(add_form.image_url.data, 500),
    )
    db.session.add(item)
    db.session.flush()

    err = _save_cover_file(item, add_form.cover_file.data)
    if err:
        db.session.rollback()
        flask.flash(err, "danger")
        return flask.redirect(flask.url_for("jeux.admin"))

    db.session.commit()
    flask.flash(
        _("Le jeu « %(nom)s » a été ajouté au catalogue.", nom=item.nom),
        "success",
    )
    return flask.redirect(flask.url_for("jeux.admin"))


@bp.route("/admin/upload", methods=["POST"])
@context.permission_only(PermissionType.write, PermissionScope.jeux)
def admin_upload() -> typing.RouteReturn:
    """Handle ODS/XLSX catalogue upload and upsert into DB."""
    form = UploadCatalogueForm()
    if not form.validate_on_submit():
        _flash_form_errors(form)
        return flask.redirect(flask.url_for("jeux.admin"))

    file_data = form.fichier.data
    filename = file_data.filename or "catalogue.ods"
    stream = io.BytesIO(file_data.read())

    try:
        added, updated = _import_catalogue(stream, filename)
    except Exception as exc:
        flask.flash(
            _("Erreur lors de l'import : %(err)s", err=str(exc)), "danger"
        )
        return flask.redirect(flask.url_for("jeux.admin"))

    flask.flash(
        _(
            "Import réussi : %(a)d nouvelle(s) entrée(s) ajoutée(s), %(u)d mise(s) à jour.",
            a=added,
            u=updated,
        ),
        "success",
    )
    return flask.redirect(flask.url_for("jeux.admin"))


@bp.route("/admin/export", methods=["GET"])
@context.permission_only(PermissionType.write, PermissionScope.jeux)
def admin_export() -> typing.RouteReturn:
    """Export the current Jeux database as an ODS file."""
    items = db.session.scalars(
        sa.select(JeuItem).order_by(
            sa.func.lower(JeuItem.categorie),
            sa.func.lower(JeuItem.nom),
        )
    ).all()
    ods_stream = _export_catalogue_ods(list(items))
    today_str = datetime.date.today().strftime("%Y-%m-%d")
    return flask.send_file(
        ods_stream,
        mimetype="application/vnd.oasis.opendocument.spreadsheet",
        as_attachment=True,
        download_name=f"catalogue_jeux_{today_str}.ods",
    )


@bp.route("/edit/<int:item_id>", methods=["POST"])
@context.permission_only(PermissionType.write, PermissionScope.jeux)
def edit_item(item_id: int) -> typing.RouteReturn:
    """Update Jeu item details (metadata, cover image file or URL, resume, nota_bene)."""
    item = db.session.get(JeuItem, item_id)
    if not item:
        flask.flash(_("Entrée introuvable."), "danger")
        return flask.redirect(flask.url_for("jeux.main"))

    form = flask.request.form
    if nom := _clean_str(form.get("nom"), 500):
        item.nom = nom
    if categorie := _clean_str(form.get("categorie"), 100):
        item.categorie = categorie
    item.temps_de_jeu = _clean_str(form.get("temps_de_jeu"), 100)
    item.nb_joueurs = _clean_str(form.get("nb_joueurs"), 100)
    item.popularite = _clean_str(form.get("popularite"), 100)

    if form.get("remove_cover") == "1":
        _remove_cover_file(item)

    err = _save_cover_file(item, flask.request.files.get("cover_file"))
    if err:
        flask.flash(err, "danger")
        return _redirect_to_main()

    image_url = _clean_str(form.get("image_url"), 500)
    if image_url:
        item.image_url = image_url
    elif item.image_url is not None:
        item.image_url = ""

    resume = _clean_str(form.get("resume"))
    if resume:
        item.resume = resume
    elif item.resume is not None:
        item.resume = ""

    nota_bene = _clean_str(form.get("nota_bene"))
    item.nota_bene = nota_bene
    if not nota_bene:
        item.commentaires = None

    db.session.commit()

    flask.flash(
        _("La fiche « %(nom)s » a été mise à jour.", nom=item.nom),
        "success",
    )

    return _redirect_to_main()


@bp.route("/delete/<int:item_id>", methods=["POST"])
@context.permission_only(PermissionType.write, PermissionScope.jeux)
def delete_item(item_id: int) -> typing.RouteReturn:
    """Delete a Jeu item."""
    item = db.session.get(JeuItem, item_id)
    if not item:
        flask.flash(_("Entrée introuvable."), "danger")
        return flask.redirect(flask.url_for("jeux.main"))

    nom = item.nom
    _remove_cover_file(item)

    db.session.delete(item)
    db.session.commit()

    flask.flash(
        _("La fiche « %(nom)s » a été supprimée.", nom=nom),
        "success",
    )

    return _redirect_to_main()
